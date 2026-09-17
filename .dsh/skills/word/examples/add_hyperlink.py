#!/usr/bin/env python3
"""Add a clickable hyperlink to a .docx via OOXML (python-docx has no helper).

Usage:
    python3 examples/add_hyperlink.py input.docx output.docx "https://example.com" "站点链接" --match "正文"
"""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def add_hyperlink(paragraph, url: str, text: str, color: str = "0563C1", underline: bool = True):
    """Append a w:hyperlink run to the given paragraph."""
    part = paragraph.part
    r_id = part.relate_to(url, RT.HYPERLINK, is_external=True)

    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    c = OxmlElement("w:color")
    c.set(qn("w:val"), color)
    rPr.append(c)
    if underline:
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        rPr.append(u)
    run.append(rPr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    return hyperlink


def main() -> None:
    ap = argparse.ArgumentParser(description="Add a hyperlink to a .docx")
    ap.add_argument("input", type=Path)
    ap.add_argument("output", type=Path)
    ap.add_argument("url")
    ap.add_argument("text")
    ap.add_argument("--match", type=str, default=None, help="paragraph to attach to (first match); default = new paragraph")
    args = ap.parse_args()

    doc = Document(str(args.input))

    if args.match:
        target = None
        for p in doc.paragraphs:
            if args.match in p.text:
                target = p
                break
        if target is None:
            raise SystemExit(f"no paragraph contains {args.match!r}")
        add_hyperlink(target, args.url, args.text)
    else:
        para = doc.add_paragraph()
        add_hyperlink(para, args.url, args.text)

    doc.save(str(args.output))
    print(f"saved {args.output}")


if __name__ == "__main__":
    main()
