#!/usr/bin/env python3
"""Insert content into a .docx: heading, paragraph, page break, picture, table.

Items are appended at the end, or inserted before the first paragraph matching
--before. The picture path is optional.

Usage:
    python3 examples/insert_content.py input.docx output.docx \
        --heading "结论" --heading-level 1 --paragraph "这是新加的一段。"
    python3 examples/insert_content.py input.docx output.docx \
        --picture ./figure.png --picture-width-cm 8 --before "第一章"
    python3 examples/insert_content.py input.docx output.docx \
        --table 'a,b|1,2'
"""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Cm


def find_before(doc: Document, text: str):
    for p in doc.paragraphs:
        if text in p.text:
            return p
    return None


def insert_heading(doc: Document, anchor, text: str, level: int):
    if anchor is None:
        doc.add_heading(text, level=level)
    else:
        p = anchor.insert_paragraph_before(text, style=f"Heading {level}")
        p.style = doc.styles[f"Heading {level}"]


def insert_paragraph(doc: Document, anchor, text: str):
    if anchor is None:
        doc.add_paragraph(text)
    else:
        anchor.insert_paragraph_before(text)


def insert_page_break(doc: Document, anchor):
    p = doc.add_paragraph() if anchor is None else anchor.insert_paragraph_before()
    run = p.add_run()
    run.add_break(WD_BREAK.PAGE)


def insert_picture(doc: Document, anchor, path: str, width_cm):
    if anchor is None:
        doc.add_picture(path, width=Cm(width_cm))
    else:
        anchor.insert_paragraph_before().add_run().add_picture(path, width=Cm(width_cm))


def insert_table(doc: Document, anchor, csv: str):
    rows = [r.split(",") for r in csv.split("|")]
    n_cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=n_cols) if anchor is None else None
    if anchor is not None:
        table = doc.add_table(rows=len(rows), cols=n_cols)
        anchor._p.addprevious(table._tbl)
    for ri, row in enumerate(rows):
        for ci in range(n_cols):
            table.rows[ri].cells[ci].text = row[ci] if ci < len(row) else ""
    table.style = doc.styles["Table Grid"] if "Table Grid" in [s.name for s in doc.styles] else table.style


def main() -> None:
    ap = argparse.ArgumentParser(description="Insert content into a .docx")
    ap.add_argument("input", type=Path)
    ap.add_argument("output", type=Path)
    ap.add_argument("--heading", type=str, default=None)
    ap.add_argument("--heading-level", type=int, default=1)
    ap.add_argument("--paragraph", type=str, default=None)
    ap.add_argument("--page-break", action="store_true")
    ap.add_argument("--picture", type=str, default=None)
    ap.add_argument("--picture-width-cm", type=float, default=8.0)
    ap.add_argument("--table", type=str, default=None, help="CSV rows separated by '|'")
    ap.add_argument("--before", type=str, default=None, help="insert before first paragraph containing this")
    args = ap.parse_args()

    doc = Document(str(args.input))
    anchor = find_before(doc, args.before) if args.before else None

    if args.heading:
        insert_heading(doc, anchor, args.heading, args.heading_level)
    if args.paragraph:
        insert_paragraph(doc, anchor, args.paragraph)
    if args.page_break:
        insert_page_break(doc, anchor)
    if args.picture:
        insert_picture(doc, anchor, args.picture, args.picture_width_cm)
    if args.table:
        insert_table(doc, anchor, args.table)

    doc.save(str(args.output))
    print(f"saved {args.output}")


if __name__ == "__main__":
    main()
