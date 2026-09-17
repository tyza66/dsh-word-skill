#!/usr/bin/env python3
"""Audit heading hierarchy, numbering, and direct formatting in a DOCX.

Generic QA tool for word skill. Surfaces the most common causes of
"why does this paragraph look different" issues before rendering:

- Heading styles used and level jumps (H1 -> H3 without H2)
- Numbering on paragraphs that lack a Heading style (breaks TOC)
- Direct (run-level) formatting overrides on paragraphs and tables
- Font usage summary

Does not mutate the document.
"""

from __future__ import annotations

import argparse
from collections import Counter
from pathlib import Path

from docx import Document


def _heading_level(style_name):
    if not style_name:
        return None
    s = style_name.strip().lower()
    if not s.startswith("heading"):
        return None
    parts = s.split()
    if len(parts) < 2:
        return None
    try:
        return int(parts[1])
    except Exception:
        return None


def _has_numbering(p) -> bool:
    pPr = p._p.pPr
    if pPr is None:
        return False
    return pPr.numPr is not None


def _has_direct_run_formatting(run) -> bool:
    f = run.font
    return any(
        v is not None
        for v in [
            run.bold,
            run.italic,
            run.underline,
            f.name,
            f.size,
            f.color.rgb if f.color else None,
        ]
    )


def _iter_paragraphs(doc: Document):
    for p in doc.paragraphs:
        yield p
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    yield p
    for section in doc.sections:
        for p in section.header.paragraphs:
            yield p
        for p in section.footer.paragraphs:
            yield p


def main() -> None:
    ap = argparse.ArgumentParser(description="Audit heading hierarchy and direct formatting")
    ap.add_argument("docx", type=Path)
    ap.add_argument("--max_findings", type=int, default=20)
    args = ap.parse_args()

    if not args.docx.exists():
        raise FileNotFoundError(args.docx)

    doc = Document(str(args.docx))
    h_counts: Counter[int] = Counter()
    jumps: list[str] = []
    numbered_non_heading: list[str] = []
    direct_overrides: list[str] = []
    font_names: Counter[str] = Counter()
    font_sizes: Counter[float] = Counter()

    last_h: int | None = None

    for i, p in enumerate(_iter_paragraphs(doc), start=1):
        style = getattr(p.style, "name", None)
        lvl = _heading_level(style)
        if lvl is not None:
            h_counts[lvl] += 1
            if last_h is not None and lvl > last_h + 1:
                jumps.append(f"p#{i}: Heading {last_h} -> Heading {lvl}: {p.text[:80]!r}")
            last_h = lvl
        if _has_numbering(p) and lvl is None:
            txt = (p.text or "").strip()
            if txt:
                numbered_non_heading.append(f"p#{i}: style={style!r} text={txt[:80]!r}")
        for run in p.runs:
            if _has_direct_run_formatting(run):
                txt = (p.text or "").strip()
                if txt:
                    direct_overrides.append(f"p#{i}: run override in {txt[:60]!r}")
            if run.font.name:
                font_names[run.font.name] += 1
            if run.font.size:
                font_sizes[round(run.font.size.pt, 1)] += 1

    print("HEADING STYLE COUNTS")
    if not h_counts:
        print("- (no Heading styles found)")
    else:
        for lvl in sorted(h_counts):
            print(f"- Heading {lvl}: {h_counts[lvl]}")

    if jumps:
        print("\nHEADING LEVEL JUMPS (review)")
        for s in jumps[: args.max_findings]:
            print(f"- {s}")
        if len(jumps) > args.max_findings:
            print(f"- ... ({len(jumps) - args.max_findings} more)")

    if numbered_non_heading:
        print("\nNUMBERING WITHOUT HEADING STYLE (TOC risk)")
        for s in numbered_non_heading[: args.max_findings]:
            print(f"- {s}")

    if direct_overrides:
        print("\nDIRECT RUN-LEVEL FORMATTING (review)")
        for s in direct_overrides[: args.max_findings]:
            print(f"- {s}")

    print("\nFONT USAGE")
    for name, count in font_names.most_common(20):
        print(f"- {name}: {count} runs")
    for size, count in font_sizes.most_common(10):
        print(f"- {size}pt: {count} runs")


if __name__ == "__main__":
    main()
